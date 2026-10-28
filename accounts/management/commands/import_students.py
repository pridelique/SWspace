import re
from collections import Counter
from pathlib import Path

from django.conf import settings
from django.contrib.auth import get_user_model
from django.core.management.base import BaseCommand, CommandError

User = get_user_model()

ROOM_RE = re.compile(r'ปีที่\s*(\d)\.(\d+)')
STUDENT_ID_RE = re.compile(r'\d{5}')
SKIP_SHEETS = {'สรุป'}


def parse_roster(path):
    """Yield one dict per student row of the roster workbook.

    The classroom is read from each block's title ("...ปีที่ 5.7...") as text,
    never from a number, so 1.10 can't collapse into 1.1.
    """
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=True)
    for ws in wb:
        if ws.title in SKIP_SHEETS:
            continue
        classroom = None
        for row in ws.iter_rows(values_only=True):
            for cell in row[:3]:
                if isinstance(cell, str):
                    m = ROOM_RE.search(cell)
                    if m:
                        classroom = f'{m.group(1)}.{m.group(2)}'
            number, raw_id = row[0], row[1]
            if not isinstance(number, int) or raw_id is None:
                continue
            student_id = str(raw_id).strip()
            if not STUDENT_ID_RE.fullmatch(student_id) or not classroom:
                continue
            prefix, first, last = (str(c).strip() if c else '' for c in row[2:5])
            notes = [str(c).strip() for c in row[5:] if isinstance(c, str)]
            yield {
                'student_id': student_id,
                'classroom': classroom,
                'full_name': f'{prefix}{first} {last}'.strip(),
                'left_school': any('ออก' in n for n in notes),
            }


class Command(BaseCommand):
    help = (
        'Create/update student accounts from the roster .xlsx. '
        'Login: <student_id>@satriwit.ac.th, password: <student_id>_<grade.room>.'
    )

    def add_arguments(self, parser):
        parser.add_argument('path', help='Path to the roster .xlsx')
        parser.add_argument('--dry-run', action='store_true', help='Print counts only; write nothing')
        parser.add_argument(
            '--include-left', action='store_true',
            help='Also create accounts for rows marked "ออก" (left school). Skipped by default.',
        )

    def handle(self, *args, **opts):
        path = Path(opts['path'])
        if not path.is_file():
            raise CommandError(f'File not found: {path}')

        students = list(parse_roster(path))
        ids = Counter(s['student_id'] for s in students)
        dups = sorted(i for i, n in ids.items() if n > 1)
        if dups:
            raise CommandError(f'Duplicate student IDs in roster: {len(dups)} (aborting)')

        left = [s for s in students if s['left_school']]
        if not opts['include_left']:
            students = [s for s in students if not s['left_school']]

        by_grade = Counter(s['classroom'].split('.')[0] for s in students)
        self.stdout.write(f'Rows to import: {len(students)}  (skipped as left school: '
                          f'{0 if opts["include_left"] else len(left)})')
        self.stdout.write('By grade: ' + ', '.join(f'ม.{g}={n}' for g, n in sorted(by_grade.items())))
        if opts['dry_run']:
            self.stdout.write('Dry run - nothing written.')
            return

        domain = settings.STUDENT_EMAIL_DOMAIN
        created = updated = unchanged = 0
        protected = []
        total = len(students)
        for i, s in enumerate(students, 1):
            email = f"{s['student_id']}{domain}"
            password = f"{s['student_id']}_{s['classroom']}"
            grade_level = f"ม.{s['classroom'].split('.')[0]}"
            user = User.objects.filter(email=email).first()

            if user is None:
                User.objects.create_user(
                    email=email, password=password, full_name=s['full_name'], role='student',
                    student_id=s['student_id'], classroom=s['classroom'], grade_level=grade_level,
                )
                created += 1
            elif user.role == 'committee' or user.is_staff or user.is_superuser:
                protected.append(email)
            else:
                # Hash only when needed: first time under the roster scheme, or room changed.
                needs_password = user.student_id != s['student_id'] or user.classroom != s['classroom']
                changed = needs_password or user.full_name != s['full_name'] or user.grade_level != grade_level
                if changed:
                    user.student_id = s['student_id']
                    user.classroom = s['classroom']
                    user.grade_level = grade_level
                    user.full_name = s['full_name']
                    if needs_password:
                        user.set_password(password)
                    user.save()
                    updated += 1
                else:
                    unchanged += 1

            if i % 250 == 0:
                self.stdout.write(f'  {i}/{total}')

        self.stdout.write(self.style.SUCCESS(
            f'Done. created={created} updated={updated} unchanged={unchanged} protected={len(protected)}'
        ))
        if protected:
            self.stdout.write(self.style.WARNING(
                'Committee/staff accounts matching roster emails were NOT touched:'))
            for e in protected:
                self.stdout.write(f'  {e}')
